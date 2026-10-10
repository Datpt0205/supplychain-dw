"""Unit: versioned document templates and the in-process DOCX renderer
(supply-chain ticket ai-automation/03).

Measured on real files python-docx writes and reads: a value is data (a
placeholder spelled inside a value, a formula, a field code are printed as
written), a template is checked when loaded, and a tenant's override never
reaches another tenant.
"""

from __future__ import annotations

import io
import uuid
import zipfile

import pytest
import yaml
from docx import Document

from dw_agent_runtime.adapters.docx_templates import DocxRenderer, DocxTemplateInspector
from dw_agent_runtime.doc_templates import DocTemplateRegistry, TemplateValue
from dw_kernel.errors import ConfigError, NotFoundError

pytestmark = pytest.mark.unit

TENANT_A, TENANT_B = uuid.uuid4(), uuid.uuid4()


def _spec(version: str = "1.0.0", title: str = "Đơn đặt hàng") -> bytes:
    return yaml.safe_dump(
        {
            "schema_version": "1.0",
            "template_id": "test.purchase_order",
            "version": version,
            "title": title,
            "doc_type": "purchase_order",
            "fields": [
                {"name": "po_reference", "label": "Số PO", "required": True},
                {"name": "notes", "label": "Ghi chú"},
                {
                    "name": "lines",
                    "label": "Dòng hàng",
                    "kind": "table",
                    "columns": [
                        {"name": "sku", "label": "SKU"},
                        {"name": "quantity", "label": "Số lượng", "kind": "number"},
                    ],
                },
            ],
        },
        allow_unicode=True,
    ).encode("utf-8")


def _docx(heading: str = "ĐƠN ĐẶT HÀNG", extra: str | None = None) -> bytes:
    document = Document()
    document.add_paragraph(heading)
    document.add_paragraph("Số PO: {{po_reference}}")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "SKU"
    table.cell(0, 1).text = "Số lượng"
    table.cell(1, 0).text = "{{lines.sku}}"
    table.cell(1, 1).text = "{{lines.quantity}}"
    document.add_paragraph("Ghi chú: {{notes}}")
    if extra is not None:
        document.add_paragraph(extra)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _registry() -> DocTemplateRegistry:
    return DocTemplateRegistry(inspector=DocxTemplateInspector())


def _text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def test_a_template_renders_its_fields_and_repeats_its_table_row() -> None:
    template = _registry().load_bytes(_spec(), _docx())
    rendered = DocxRenderer().render(
        template,
        {
            "po_reference": "PO-0012",
            "lines": [{"sku": "NOI-24", "quantity": "100"}, {"sku": "CHAO-28", "quantity": None}],
        },
    )
    text = _text(rendered.data)
    assert "Số PO: PO-0012" in text
    assert "NOI-24 | 100" in text
    # A gap is visible in the document, never blank.
    assert "CHAO-28 | [cần điền: Dòng hàng / Số lượng]" in text
    assert "Ghi chú: [cần điền: Ghi chú]" in text
    assert "{{" not in text


def test_a_value_is_data_never_a_placeholder_a_formula_or_a_field_code() -> None:
    template = _registry().load_bytes(_spec(), _docx())
    hostile: dict[str, TemplateValue] = {
        "po_reference": "{{notes}}",
        "notes": '=HYPERLINK("http://evil.example","x") <w:fldSimple w:instr="DDE"/>\x07',
        "lines": [{"sku": "{{po_reference}}", "quantity": "=1+1"}],
    }
    rendered = DocxRenderer().render(template, hostile)
    text = _text(rendered.data)
    assert "Số PO: {{notes}}" in text
    assert "{{po_reference}} | =1+1" in text
    assert '=HYPERLINK("http://evil.example","x")' in text
    xml = _xml(rendered.data)
    assert "<w:fldSimple" not in xml and "w:instrText" not in xml
    assert "\x07" not in text
    with zipfile.ZipFile(io.BytesIO(rendered.data)) as archive:
        assert not any("vbaProject" in name for name in archive.namelist())


def test_a_template_whose_declaration_and_placeholders_differ_is_refused_by_name() -> None:
    with pytest.raises(ConfigError, match=r"test.purchase_order@1.0.0"):
        _registry().load_bytes(_spec(), _docx(extra="Hạn: {{due_date}}"))


def test_a_placeholder_split_across_runs_is_refused() -> None:
    document = Document(io.BytesIO(_docx()))
    paragraph = document.add_paragraph("Hạn: {{po_")
    paragraph.add_run("reference}}")
    out = io.BytesIO()
    document.save(out)
    with pytest.raises(ConfigError, match="split"):
        DocxTemplateInspector().placeholders(out.getvalue())


def test_a_template_carrying_macros_is_refused() -> None:
    source = io.BytesIO(_docx())
    out = io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(out, "w") as copy:
        for item in original.infolist():
            copy.writestr(item, original.read(item.filename))
        copy.writestr("word/vbaProject.bin", b"\x00macro")
    with pytest.raises(ConfigError, match="macros"):
        DocxTemplateInspector().placeholders(out.getvalue())


def test_a_tenants_override_reaches_only_that_tenant_and_others_fall_back() -> None:
    registry = _registry()
    registry.load_bytes(_spec(), _docx("ĐƠN ĐẶT HÀNG (mẫu chung)"))
    registry.load_bytes(
        _spec(title="PO của A"), _docx("ĐƠN ĐẶT HÀNG CÔNG TY A"), tenant_id=TENANT_A
    )
    own = registry.resolve("test.purchase_order", "1.0.0", tenant_id=TENANT_A)
    other = registry.resolve("test.purchase_order", "1.0.0", tenant_id=TENANT_B)
    platform = registry.resolve("test.purchase_order", "1.0.0")
    assert own.spec.title == "PO của A"
    assert other.checksum == platform.checksum != own.checksum
    assert "CÔNG TY A" not in _text(DocxRenderer().render(other, {}).data)
    # A version only A holds is not B's.
    registry.load_bytes(_spec(version="2.0.0"), _docx(), tenant_id=TENANT_A)
    with pytest.raises(NotFoundError):
        registry.resolve("test.purchase_order", "2.0.0", tenant_id=TENANT_B)
    assert [t.spec.ref for t in registry.platform_templates()] == ["test.purchase_order@1.0.0"]


def test_one_version_twice_in_one_layer_is_refused() -> None:
    registry = _registry()
    registry.load_bytes(_spec(), _docx())
    with pytest.raises(ConfigError, match="already registered"):
        registry.load_bytes(_spec(), _docx())
