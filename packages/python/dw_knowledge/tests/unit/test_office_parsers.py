"""Unit: in-process text of PDF, DOCX, XLSX and EML, measured on real files the
pinned libraries write (failure-modes #4: a library is run, not trusted)."""

from __future__ import annotations

import io
from email.message import EmailMessage

import pytest
from docx import Document
from openpyxl import Workbook

from dw_knowledge.adapters.office_parsers import (
    DOCX,
    EML,
    PDF,
    XLSX,
    InProcessDocumentParser,
)

pytestmark = pytest.mark.unit


def _pdf(text: str) -> bytes:
    """A one-page PDF with a text layer, written by hand: enough for pypdf."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Biên bản đánh giá mẫu")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tiêu chí"
    table.cell(0, 1).text = "Kết quả"
    table.cell(1, 0).text = "Độ dày đáy"
    table.cell(1, 1).text = "Không đạt"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Báo giá"
    sheet.append(["SKU", "Số lượng", "Đơn giá"])
    sheet.append(["NOI-24", 100, 2.5])
    # A formula with no cached value: never evaluated, read as empty.
    sheet["D2"] = "=B2*C2"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


async def test_a_pdf_text_layer_is_read_page_by_page() -> None:
    parsed = await InProcessDocumentParser().parse(_pdf("Unit price USD 2.50"), PDF, "q.pdf")
    assert "Unit price USD 2.50" in parsed.text
    assert parsed.text.startswith("--- Trang 1 ---")


async def test_docx_paragraphs_and_table_cells_are_read() -> None:
    parsed = await InProcessDocumentParser().parse(_docx(), DOCX, "bb.docx")
    assert "Biên bản đánh giá mẫu" in parsed.text
    assert "Độ dày đáy | Không đạt" in parsed.text


async def test_xlsx_values_are_read_and_a_formula_is_never_evaluated() -> None:
    parsed = await InProcessDocumentParser().parse(_xlsx(), XLSX, "q.xlsx")
    assert "## Báo giá" in parsed.text
    assert "NOI-24 | 100 | 2.5" in parsed.text
    assert "=B2*C2" not in parsed.text and "250" not in parsed.text


async def test_an_email_body_and_its_readable_attachments_are_read() -> None:
    message = EmailMessage()
    message["From"] = "ncc@example.com"
    message["To"] = "cung-ung@example.com"
    message["Subject"] = "Xác nhận đơn hàng"
    message.set_content("Chúng tôi xác nhận số lượng 1000 cái.")
    message.add_attachment(
        _docx(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="bb.docx",
    )
    message.add_attachment(b"\x89PNG", maintype="image", subtype="png", filename="anh.png")
    parsed = await InProcessDocumentParser().parse(bytes(message), EML, "x.eml")
    assert "Subject: Xác nhận đơn hàng" in parsed.text
    assert "Chúng tôi xác nhận số lượng 1000 cái." in parsed.text
    assert "--- Tệp đính kèm: bb.docx ---" in parsed.text
    assert "Độ dày đáy | Không đạt" in parsed.text
    assert "[tệp đính kèm chưa đọc: anh.png (image/png)]" in parsed.text


async def test_a_long_text_is_cut_and_says_so() -> None:
    parsed = await InProcessDocumentParser(max_chars=10).parse(_docx(), DOCX, "bb.docx")
    assert len(parsed.text) == 10
    assert parsed.warnings == ("truncated",)


def test_only_the_four_types_are_supported() -> None:
    parser = InProcessDocumentParser()
    assert all(parser.supports(t, "f") for t in (PDF, DOCX, XLSX, EML))
    assert not parser.supports("image/png", "a.png")
    assert not parser.supports("application/vnd.ms-outlook", "a.msg")
