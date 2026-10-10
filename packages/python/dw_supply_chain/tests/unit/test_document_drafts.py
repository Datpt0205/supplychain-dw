"""Unit: document drafts and the tenant's templates (ADR 0025 point 4; ticket
ai-automation/03).

The templates are the SHIPPED ones (`configs/doc_templates`), loaded through the
real registry and inspector, and rendered by the real DOCX renderer. Fakes hold
the rows and honour RLS the way the database does.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Mapping
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, cast

import pytest
import yaml
from docx import Document

from dw_agent_runtime.adapters.docx_templates import DocxRenderer, DocxTemplateInspector
from dw_agent_runtime.doc_templates import DocTemplateRegistry
from dw_kernel.errors import ConflictError, DomainError, NotFoundError, PermissionDeniedError
from dw_kernel.ports import FixedClock, Uuid4Generator
from dw_platform.application.access_context import AccessContext
from dw_platform.application.authorization import ScopeAuthorizationService
from dw_platform.domain.audit import AuditEvent
from dw_supply_chain.application.document_drafts import (
    HIDDEN,
    FieldInput,
    GetDocumentDraft,
    ListCaseDrafts,
    NewDocumentDraft,
    NewDraftDecision,
    PrepareDocumentDraft,
    RejectDocumentDraft,
    RenderDocumentDraft,
    ReviseDocumentDraft,
    SetDocTemplateOverride,
    StoredTemplate,
    TenantDocTemplates,
)
from dw_supply_chain.application.handlers import (
    ACTION_DUTIES_WRITE,
    COMMERCIAL_READ,
    COMMERCIAL_WRITE,
    DOCUMENT_READ,
    DOCUMENT_WRITE,
)
from dw_supply_chain.domain.case_document import CaseKind, DocumentType
from dw_supply_chain.domain.document_draft import (
    DocumentDraft,
    DraftDecision,
    DraftSource,
    DraftStatus,
)

pytestmark = pytest.mark.unit

REPO_ROOT = Path(__file__).resolve().parents[5]
TENANT, WORKSPACE, OTHER_TENANT, OTHER_WORKSPACE = (uuid.uuid4() for _ in range(4))
NOW = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
AUTHZ = ScopeAuthorizationService()
IDS = Uuid4Generator()
CLOCK = FixedClock(NOW)
READER = frozenset({DOCUMENT_READ})
WRITER = frozenset({DOCUMENT_READ, DOCUMENT_WRITE})
PRICED = WRITER | {COMMERCIAL_READ, COMMERCIAL_WRITE}


def _context(
    scopes: frozenset[str], *, tenant: uuid.UUID = TENANT, workspace: uuid.UUID = WORKSPACE
) -> AccessContext:
    return AccessContext(
        tenant_id=tenant,
        workspace_id=workspace,
        principal_id=uuid.uuid4(),
        roles=frozenset({"member"}),
        scopes=scopes,
        plan_id="professional",
    )


def _platform_registry() -> DocTemplateRegistry:
    registry = DocTemplateRegistry(inspector=DocxTemplateInspector())
    registry.load_directory(REPO_ROOT / "configs" / "doc_templates")
    return registry


class FakeCases:
    def __init__(self) -> None:
        self.cases: dict[uuid.UUID, tuple[uuid.UUID, uuid.UUID]] = {}

    def add(self, tenant: uuid.UUID = TENANT, workspace: uuid.UUID = WORKSPACE) -> uuid.UUID:
        case_id = uuid.uuid4()
        self.cases[case_id] = (tenant, workspace)
        return case_id

    async def case_workspace(self, context: AccessContext, case_id: uuid.UUID) -> uuid.UUID | None:
        found = self.cases.get(case_id)
        return None if found is None or found[0] != context.tenant_id else found[1]


@dataclass
class FakeDrafts:
    rows: list[DocumentDraft] = field(default_factory=list)
    decisions: dict[uuid.UUID, tuple[DraftDecision, str | None]] = field(default_factory=dict)
    audits: list[AuditEvent] = field(default_factory=list)

    def _sees(self, context: AccessContext, d: DocumentDraft) -> bool:
        return (d.tenant_id, d.workspace_id) == (context.tenant_id, context.workspace_id)

    def _with_state(self, d: DocumentDraft) -> DocumentDraft:
        latest = max(r.version for r in self.rows if r.lineage_id == d.lineage_id)
        decision = self.decisions.get(d.id)
        return replace(
            d,
            latest_version=latest,
            decision=None if decision is None else decision[0],
            decision_reason=None if decision is None else decision[1],
        )

    async def add(
        self, context: AccessContext, draft: NewDocumentDraft, *, audit: AuditEvent
    ) -> DocumentDraft:
        if any(r.lineage_id == draft.lineage_id and r.version == draft.version for r in self.rows):
            raise ConflictError("version taken")
        row = DocumentDraft(
            id=draft.id,
            tenant_id=context.tenant_id,
            workspace_id=context.workspace_id,
            lineage_id=draft.lineage_id,
            version=draft.version,
            case_kind=draft.case_kind,
            case_id=draft.case_id,
            doc_type=draft.doc_type,
            template_id=draft.template_id,
            template_version=draft.template_version,
            prompt_id=draft.prompt_id,
            prompt_version=draft.prompt_version,
            fields=draft.fields,
            gaps=tuple(draft.gaps),
            sources=tuple(
                DraftSource(document_id=uuid.UUID(str(s["document_id"])), sha256=str(s["sha256"]))
                for s in draft.sources
            ),
            content_sha256=draft.content_sha256,
            created_by=context.principal_id,
            created_at=NOW,
        )
        self.rows.append(row)
        self.audits.append(audit)
        return self._with_state(row)

    async def get(self, context: AccessContext, draft_id: uuid.UUID) -> DocumentDraft | None:
        for row in self.rows:
            if row.id == draft_id and self._sees(context, row):
                return self._with_state(row)
        return None

    async def latest_for_case(
        self, context: AccessContext, case_kind: CaseKind, case_id: uuid.UUID
    ) -> list[DocumentDraft]:
        mine = [
            self._with_state(r)
            for r in self.rows
            if self._sees(context, r) and r.case_id == case_id
        ]
        return [r for r in mine if r.version == r.latest_version]

    async def decide(
        self, context: AccessContext, decision: NewDraftDecision, *, audit: AuditEvent
    ) -> None:
        if decision.draft_id in self.decisions:
            raise ConflictError("already decided")
        self.decisions[decision.draft_id] = (decision.decision, decision.reason)
        self.audits.append(audit)


@dataclass
class FakeOverrides:
    rows: dict[tuple[uuid.UUID, str, str], StoredTemplate] = field(default_factory=dict)

    async def get(
        self, context: AccessContext, template_id: str, version: str
    ) -> StoredTemplate | None:
        return self.rows.get((context.tenant_id, template_id, version))

    async def list_refs(self, context: AccessContext) -> list[tuple[str, str]]:
        return [(t, v) for (tenant, t, v) in self.rows if tenant == context.tenant_id]

    async def add(
        self, context: AccessContext, template: StoredTemplate, *, audit: AuditEvent
    ) -> None:
        key = (context.tenant_id, template.template_id, template.version)
        if key in self.rows:
            raise ConflictError("exists")
        self.rows[key] = template


@dataclass
class World:
    cases: FakeCases = field(default_factory=FakeCases)
    drafts: FakeDrafts = field(default_factory=FakeDrafts)
    overrides: FakeOverrides = field(default_factory=FakeOverrides)
    registry: DocTemplateRegistry = field(default_factory=_platform_registry)

    @property
    def templates(self) -> TenantDocTemplates:
        return TenantDocTemplates(registry=self.registry, overrides=self.overrides)

    def lookups(self) -> dict[CaseKind, FakeCases]:
        return {CaseKind.PO: self.cases, CaseKind.PRODUCT: self.cases}

    async def prepare(
        self,
        values: Mapping[str, Any],
        *,
        case_id: uuid.UUID | None = None,
        doc_type: DocumentType = DocumentType.BOD_SUBMISSION,
        context: AccessContext | None = None,
    ) -> DocumentDraft:
        return await PrepareDocumentDraft(
            cases=self.lookups(),
            drafts=self.drafts,
            templates=self.templates,
            ids=IDS,
            clock=CLOCK,
        ).handle(
            context or _context(frozenset()),
            case_kind=CaseKind.PRODUCT,
            case_id=case_id or self.cases.add(),
            doc_type=doc_type,
            values={n: FieldInput(value=v) for n, v in values.items()},
        )

    def revise(self) -> ReviseDocumentDraft:
        return ReviseDocumentDraft(
            drafts=self.drafts, templates=self.templates, authz=AUTHZ, ids=IDS, clock=CLOCK
        )

    def get(self) -> GetDocumentDraft:
        return GetDocumentDraft(drafts=self.drafts, templates=self.templates, authz=AUTHZ)

    def render(self) -> RenderDocumentDraft:
        return RenderDocumentDraft(
            drafts=self.drafts, templates=self.templates, renderer=DocxRenderer(), authz=AUTHZ
        )


SUBMISSION = {
    "proposal_code": "DX-001",
    "product_name": "Nồi inox 24cm",
    "supplier_name": "Minh Phát",
    "evaluation_result": "Đạt vòng 2",
    "unit_price": "2,50",
    "currency": "USD",
}


def _text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


async def test_a_prepared_draft_has_typed_values_named_gaps_and_a_stable_hash() -> None:
    world = World()
    draft = await world.prepare(SUBMISSION)
    assert draft.version == 1 and draft.lineage_id == draft.id
    assert draft.fields["unit_price"]["value"] == "2.50"
    assert set(draft.gaps) == {"recommendation"}
    again = await world.prepare(SUBMISSION)
    assert again.content_sha256 == draft.content_sha256


async def test_a_value_the_template_does_not_declare_or_a_wrong_kind_is_a_422() -> None:
    world = World()
    with pytest.raises(DomainError) as refused:
        await world.prepare({**SUBMISSION, "bank_account": "0071000123456", "moq": "nhiều"})
    errors = cast(list[dict[str, str]], refused.value.details["errors"])
    assert {(e["field"], e["error"]) for e in errors} == {
        ("bank_account", "unknown_field"),
        ("moq", "not_a_number"),
    }


async def test_a_draft_of_another_workspace_or_tenant_is_not_found() -> None:
    world = World()
    with pytest.raises(NotFoundError):
        await world.prepare(SUBMISSION, case_id=world.cases.add(workspace=OTHER_WORKSPACE))
    draft = await world.prepare(SUBMISSION)
    for other in (
        _context(PRICED, workspace=OTHER_WORKSPACE),
        _context(PRICED, tenant=OTHER_TENANT),
    ):
        with pytest.raises(NotFoundError):
            await world.get().handle(other, draft.id)
        with pytest.raises(NotFoundError):
            await world.revise().handle(other, draft.id, values={"recommendation": "Duyệt"})


async def test_an_edit_is_a_new_version_and_the_old_one_stays_readable() -> None:
    world = World()
    first = await world.prepare(SUBMISSION)
    second = await world.revise().handle(
        _context(WRITER), first.id, values={"recommendation": "Đề nghị duyệt"}
    )
    assert (second.version, second.lineage_id) == (2, first.lineage_id)
    assert second.gaps == ()
    assert second.fields["recommendation"]["source"]["edited_by"]
    assert second.content_sha256 != first.content_sha256
    old = await world.get().handle(_context(READER), first.id)
    assert old.draft.status is DraftStatus.SUPERSEDED
    assert "recommendation" not in old.draft.fields
    # The superseded version cannot be edited or decided again.
    with pytest.raises(ConflictError):
        await world.revise().handle(_context(WRITER), first.id, values={"summary": "x"})
    with pytest.raises(ConflictError):
        await RejectDocumentDraft(
            drafts=world.drafts, templates=world.templates, authz=AUTHZ, ids=IDS, clock=CLOCK
        ).handle(_context(WRITER), first.id, reason="cũ")


async def test_a_rejection_needs_a_reason_and_closes_the_draft() -> None:
    world = World()
    draft = await world.prepare(SUBMISSION)
    reject = RejectDocumentDraft(
        drafts=world.drafts, templates=world.templates, authz=AUTHZ, ids=IDS, clock=CLOCK
    )
    with pytest.raises(DomainError):
        await reject.handle(_context(WRITER), draft.id, reason="  ")
    with pytest.raises(PermissionDeniedError):
        await reject.handle(_context(READER), draft.id, reason="sai NCC")
    rejected = await reject.handle(_context(WRITER), draft.id, reason="sai NCC")
    assert rejected.status is DraftStatus.REJECTED
    with pytest.raises(ConflictError):
        await world.revise().handle(_context(WRITER), draft.id, values={"summary": "x"})


async def test_prices_read_hidden_without_the_scope_in_the_reading_and_the_preview() -> None:
    world = World()
    draft = await world.prepare(SUBMISSION)
    blind = await world.get().handle(_context(READER), draft.id)
    assert blind.prices_visible is False
    assert blind.draft.fields["unit_price"] == {"value": None, "source": None, "redacted": True}
    preview = await world.render().handle(_context(READER), draft.id)
    text = _text(preview.file.data)
    assert HIDDEN in text and "2.50" not in text and "2,50" not in text
    seeing = await world.render().handle(_context(READER | {COMMERCIAL_READ}), draft.id)
    assert "2.50" in _text(seeing.file.data)


async def test_a_price_edit_needs_the_commercial_write_scope_and_others_keep_the_price() -> None:
    world = World()
    draft = await world.prepare(SUBMISSION)
    with pytest.raises(PermissionDeniedError):
        await world.revise().handle(_context(WRITER), draft.id, values={"unit_price": "3"})
    kept = await world.revise().handle(
        _context(WRITER), draft.id, values={"recommendation": "Duyệt"}
    )
    assert kept.fields["unit_price"]["value"] == "2.50"
    priced = await world.revise().handle(_context(PRICED), kept.id, values={"unit_price": "3"})
    assert priced.fields["unit_price"]["value"] == "3"


async def test_list_shows_the_latest_version_of_each_draft_of_the_case() -> None:
    world = World()
    case_id = world.cases.add()
    first = await world.prepare(SUBMISSION, case_id=case_id)
    await world.revise().handle(_context(WRITER), first.id, values={"summary": "x"})
    readings = await ListCaseDrafts(
        cases=world.lookups(),
        drafts=world.drafts,
        templates=world.templates,
        authz=AUTHZ,
    ).handle(_context(READER), CaseKind.PRODUCT, case_id)
    assert [r.draft.version for r in readings] == [2]


# ------------------------------------------------------------- templates --


def _override_spec(registry: DocTemplateRegistry, *, title: str, doc_type: str | None = None):
    platform = registry.resolve("supply_chain.bod_submission", "1.0.0")
    raw = platform.spec.model_dump(mode="json")
    raw["title"] = title
    if doc_type is not None:
        raw["doc_type"] = doc_type
    return yaml.safe_dump(raw, allow_unicode=True).encode("utf-8"), platform.docx


async def test_a_tenants_template_is_used_for_it_alone_and_others_fall_back() -> None:
    world = World()
    spec, docx = _override_spec(world.registry, title="TỜ TRÌNH CÔNG TY A")
    await SetDocTemplateOverride(
        registry=world.registry,
        inspector=DocxTemplateInspector(),
        overrides=world.overrides,
        authz=AUTHZ,
        ids=IDS,
        clock=CLOCK,
    ).handle(_context(frozenset({ACTION_DUTIES_WRITE})), spec=spec, docx=docx)

    mine = await world.templates.resolve(_context(READER), "supply_chain.bod_submission", "1.0.0")
    theirs = await world.templates.resolve(
        _context(READER, tenant=OTHER_TENANT), "supply_chain.bod_submission", "1.0.0"
    )
    assert mine.spec.title == "TỜ TRÌNH CÔNG TY A"
    assert theirs.spec.title == "TỜ TRÌNH BAN GIÁM ĐỐC"


async def test_a_template_override_must_vary_a_platform_template_of_the_same_type() -> None:
    world = World()
    set_override = SetDocTemplateOverride(
        registry=world.registry,
        inspector=DocxTemplateInspector(),
        overrides=world.overrides,
        authz=AUTHZ,
        ids=IDS,
        clock=CLOCK,
    )
    spec, docx = _override_spec(world.registry, title="X", doc_type="purchase_order")
    with pytest.raises(DomainError):
        await set_override.handle(_context(frozenset({ACTION_DUTIES_WRITE})), spec=spec, docx=docx)
    with pytest.raises(PermissionDeniedError):
        await set_override.handle(_context(WRITER), spec=spec, docx=docx)
    with pytest.raises(DomainError):
        await set_override.handle(
            _context(frozenset({ACTION_DUTIES_WRITE})), spec=spec, docx=b"PK\x03\x04not a docx"
        )
    assert world.overrides.rows == {}
