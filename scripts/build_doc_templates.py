"""Write the platform's neutral Supply Chain document templates.

Usage: uv run python scripts/build_doc_templates.py

Writes `configs/doc_templates/supply_chain/<template_id>@<version>.{yaml,docx}`
from the declarations below (ticket ai-automation/03). The templates are
neutral on purpose: what any company's form carries. A customer's own form
(Elmich's, QE-21) is a tenant override uploaded at run time, never a change
here. A released version is never rewritten: change a template by adding a
version (an entry with its own `version`), so a draft pinned to the old one
still renders as it did.

The .docx is written so each placeholder sits in one run (the loader refuses a
split one), and the script refuses to overwrite a file that already exists.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "configs" / "doc_templates" / "supply_chain"


def _f(name: str, label: str, kind: str = "text", required: bool = False) -> dict[str, Any]:
    field: dict[str, Any] = {"name": name, "label": label}
    if kind != "text":
        field["kind"] = kind
    if required:
        field["required"] = True
    return field


def _table(
    name: str, label: str, columns: list[tuple[str, str, str]], required: bool = False
) -> dict[str, Any]:
    field: dict[str, Any] = {
        "name": name,
        "label": label,
        "kind": "table",
        "columns": [
            {"name": c, "label": lab, **({"kind": k} if k != "text" else {})}
            for c, lab, k in columns
        ],
    }
    if required:
        field["required"] = True
    return field


TEMPLATES: list[dict[str, Any]] = [
    {
        "template_id": "supply_chain.sample_revision_request",
        "title": "PHIẾU YÊU CẦU CHỈNH SỬA MẪU",
        "doc_type": "sample_revision_request",
        "description": "Bước 4: R&D yêu cầu nhà cung cấp chỉnh sửa mẫu theo từng tiêu chí trượt.",
        "fields": [
            _f("proposal_code", "Mã đề xuất", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("sample_round", "Vòng mẫu", "number"),
            _f("requested_on", "Ngày yêu cầu", "date", required=True),
            _f("requested_by", "Người yêu cầu"),
            _table(
                "items",
                "Hạng mục chỉnh sửa",
                [
                    ("criterion", "Tiêu chí", "text"),
                    ("finding", "Hiện trạng", "text"),
                    ("requirement", "Yêu cầu chỉnh sửa", "text"),
                ],
                required=True,
            ),
            _f("due_date", "Hạn gửi mẫu chỉnh sửa", "date"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        "template_id": "supply_chain.sample_evaluation",
        "title": "BIÊN BẢN ĐÁNH GIÁ MẪU",
        "doc_type": "sample_evaluation",
        "description": "Bước 3: kết quả test mẫu theo từng tiêu chí; kết luận do người nhập.",
        "fields": [
            _f("proposal_code", "Mã đề xuất", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("sample_round", "Vòng mẫu", "number"),
            _f("evaluated_on", "Ngày đánh giá", "date", required=True),
            _f("evaluator", "Người đánh giá"),
            _table(
                "criteria",
                "Tiêu chí đánh giá",
                [
                    ("criterion", "Tiêu chí", "text"),
                    ("standard", "Chuẩn", "text"),
                    ("measured", "Đo được", "text"),
                    ("result", "Kết quả", "text"),
                ],
                required=True,
            ),
            _f("conclusion", "Kết luận", required=True),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        "template_id": "supply_chain.bod_submission",
        "title": "TỜ TRÌNH BAN GIÁM ĐỐC",
        "doc_type": "bod_submission",
        "description": "Bước 6: trình BGĐ duyệt sản phẩm sau khi mẫu đạt.",
        "fields": [
            _f("proposal_code", "Mã đề xuất", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("category", "Nhóm sản phẩm (Category)"),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("summary", "Tóm tắt sản phẩm"),
            _f("evaluation_result", "Kết quả đánh giá mẫu", required=True),
            _f("unit_price", "Đơn giá", "number"),
            _f("currency", "Tiền tệ"),
            _f("moq", "MOQ", "number"),
            _f("recommendation", "Đề xuất", required=True),
            _f("submitted_by", "Người trình"),
            _f("submitted_on", "Ngày trình", "date"),
        ],
    },
    {
        "template_id": "supply_chain.product_profile_bm04",
        "title": "HỒ SƠ SẢN PHẨM (BM04)",
        "doc_type": "product_profile_bm04",
        "description": "Bước 7: hồ sơ sản phẩm dựng từ biểu mẫu BM04 trong ứng dụng.",
        "fields": [
            _f("proposal_code", "Mã đề xuất"),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("model_code", "Mã model của NCC"),
            _f("material", "Chất liệu"),
            _f("dimensions", "Kích thước"),
            _f("net_weight_g", "Trọng lượng tịnh (g)", "number"),
            _f("capacity_ml", "Dung tích (ml)", "number"),
            _f("colours", "Màu sắc"),
            _f("packaging", "Quy cách đóng gói"),
            _f("units_per_carton", "Số sản phẩm mỗi thùng", "number"),
            _f("origin_country", "Xuất xứ"),
            _f("certifications", "Chứng nhận, tiêu chuẩn"),
            _f("unit_price", "Đơn giá", "number"),
            _f("currency", "Tiền tệ"),
            _f("moq", "MOQ", "number"),
            _f("lead_time_days", "Thời gian sản xuất (ngày)", "number"),
            _f("incoterm", "Incoterm"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Step 9 (ticket ai-automation/13): the item code and SKUs proposed by
        # code from the tenant's rule and the BM04's variants, submitted for
        # sign-off (tờ trình ký mã hàng).
        "template_id": "supply_chain.official_item_code",
        "title": "PHIẾU MÃ HÀNG CHÍNH THỨC (TRÌNH KÝ)",
        "doc_type": "official_item_code",
        "description": (
            "Bước 9: mã hàng và SKU trình ký; mã do hệ thống đề xuất theo quy tắc mã của công ty."
        ),
        "fields": [
            _f("proposal_code", "Mã đề xuất", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("category", "Nhóm sản phẩm (Category)"),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("item_code", "Mã hàng chính thức", required=True),
            _table(
                "skus",
                "SKU",
                [
                    ("sku_code", "Mã SKU", "text"),
                    ("variant_label", "Biến thể", "text"),
                    ("planned_quantity", "Số lượng dự kiến", "number"),
                ],
                required=True,
            ),
            _f("material", "Chất liệu"),
            _f("dimensions", "Kích thước"),
            _f("submitted_by", "Người trình"),
            _f("submitted_on", "Ngày trình", "date"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        "template_id": "supply_chain.purchase_order",
        "title": "ĐƠN ĐẶT HÀNG",
        "doc_type": "purchase_order",
        "description": "Bước 10: đơn đặt hàng gửi nhà cung cấp; tổng do hệ thống tính.",
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("po_date", "Ngày đặt hàng", "date", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("supplier_contact", "Người liên hệ"),
            _f("currency", "Tiền tệ", required=True),
            _f("incoterm", "Incoterm"),
            _f("payment_terms", "Điều khoản thanh toán"),
            _f("delivery_date", "Ngày giao dự kiến", "date"),
            _table(
                "lines",
                "Dòng hàng",
                [
                    ("sku_code", "SKU", "text"),
                    ("description", "Mô tả", "text"),
                    ("quantity", "Số lượng", "number"),
                    ("unit_price", "Đơn giá", "number"),
                    ("line_total", "Thành tiền", "number"),
                ],
                required=True,
            ),
            _f("order_total", "Tổng cộng", "number"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Step 10 (ticket ai-automation/14): the deposit beside the order
        # total, both computed by code; 1.0.0 stays for drafts pinned to it.
        "template_id": "supply_chain.purchase_order",
        "version": "1.1.0",
        "title": "ĐƠN ĐẶT HÀNG",
        "doc_type": "purchase_order",
        "description": (
            "Bước 10: đơn đặt hàng gửi nhà cung cấp; thành tiền, tổng và tiền cọc do hệ thống tính."
        ),
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("po_date", "Ngày đặt hàng", "date", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("supplier_contact", "Người liên hệ"),
            _f("currency", "Tiền tệ", required=True),
            _f("incoterm", "Incoterm"),
            _f("payment_terms", "Điều khoản thanh toán"),
            _f("deposit_percent", "% đặt cọc", "number"),
            _f("delivery_date", "Ngày giao dự kiến", "date"),
            _table(
                "lines",
                "Dòng hàng",
                [
                    ("sku_code", "SKU", "text"),
                    ("description", "Mô tả", "text"),
                    ("quantity", "Số lượng", "number"),
                    ("unit_price", "Đơn giá", "number"),
                    ("line_total", "Thành tiền", "number"),
                ],
                required=True,
            ),
            _f("order_total", "Tổng cộng", "number"),
            _f("deposit_amount", "Tiền đặt cọc", "number"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/15: the deposit request of step 11, every
        # amount by code (the deposit is the order total times its %).
        "template_id": "supply_chain.deposit_request",
        "title": "ĐỀ NGHỊ ĐẶT CỌC",
        "doc_type": "deposit_docs",
        "description": (
            "Bước 11: Cung ứng đề nghị Kế toán đặt cọc cho nhà cung cấp; tổng PO và tiền cọc"
            " do hệ thống tính, tài khoản thụ hưởng lấy từ danh mục NCC."
        ),
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("request_date", "Ngày đề nghị", "date", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("proforma_reference", "Số PI của NCC"),
            _f("currency", "Tiền tệ", required=True),
            _f("order_total", "Tổng giá trị PO", "number", required=True),
            _f("deposit_percent", "% đặt cọc", "number", required=True),
            _f("amount", "Số tiền đề nghị đặt cọc", "number", required=True),
            _f("bank_name", "Ngân hàng thụ hưởng"),
            _f("account_holder", "Chủ tài khoản"),
            _f("account_number", "Số tài khoản"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/15: the final payment request of step 16, the
        # balance by code (the order total less the deposit paid).
        "template_id": "supply_chain.payment_request",
        "title": "ĐỀ NGHỊ THANH TOÁN",
        "doc_type": "payment_docs",
        "description": (
            "Bước 16: Cung ứng đề nghị Kế toán thanh toán phần còn lại cho nhà cung cấp; số"
            " tiền là tổng PO trừ tiền đã cọc, do hệ thống tính."
        ),
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("request_date", "Ngày đề nghị", "date", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("invoice_reference", "Số hóa đơn của NCC"),
            _f("currency", "Tiền tệ", required=True),
            _f("order_total", "Tổng giá trị PO", "number", required=True),
            _f("deposit_paid", "Đã đặt cọc", "number", required=True),
            _f("amount", "Số tiền đề nghị thanh toán", "number", required=True),
            _f("bank_name", "Ngân hàng thụ hưởng"),
            _f("account_holder", "Chủ tài khoản"),
            _f("account_number", "Số tài khoản"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/16: the packaging content's skeleton MKT
        # completes, filled by code from the BM04; what the label must carry
        # and the BM04 does not say is a gap MKT fills.
        "template_id": "supply_chain.packaging_content",
        "title": "NỘI DUNG BAO BÌ",
        "doc_type": "packaging_content",
        "description": (
            "Bước 12: khung nội dung bao bì từ BM04 (tên, mã, SKU, chất liệu, thông số, xuất"
            " xứ); MKT điền nội dung bắt buộc còn thiếu của nhãn."
        ),
        "fields": [
            _f("po_reference", "Số PO"),
            _f("product_name", "Tên hàng hóa", required=True),
            _f("sku_codes", "Mã SKU", required=True),
            _f("material", "Chất liệu", required=True),
            _f("dimensions", "Thông số kỹ thuật", required=True),
            _f("origin", "Xuất xứ", required=True),
            _f("responsible_party", "Tổ chức chịu trách nhiệm (tên, địa chỉ)", required=True),
            _f("usage_instructions", "Hướng dẫn sử dụng, bảo quản", required=True),
            _f("warnings", "Thông tin cảnh báo", required=True),
            _f("barcode", "Mã vạch"),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/16: the user manual's skeleton from the BM04.
        "template_id": "supply_chain.user_manual",
        "title": "HƯỚNG DẪN SỬ DỤNG",
        "doc_type": "user_manual",
        "description": (
            "Bước 12: khung sách HDSD từ BM04 (tên, chất liệu, thông số); MKT viết cách dùng,"
            " bảo quản và cảnh báo."
        ),
        "fields": [
            _f("product_name", "Tên sản phẩm", required=True),
            _f("material", "Chất liệu", required=True),
            _f("dimensions", "Thông số kỹ thuật", required=True),
            _f("usage_instructions", "Cách sử dụng", required=True),
            _f("care_instructions", "Bảo quản, vệ sinh", required=True),
            _f("warnings", "Cảnh báo", required=True),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/16: Cung ứng's request to the supplier to revise
        # the colour sample; the colours are the BM04's.
        "template_id": "supply_chain.colour_revision_request",
        "title": "YÊU CẦU SỬA MẪU MÀU",
        "doc_type": "colour_revision_request",
        "description": "Bước 12: Cung ứng yêu cầu nhà cung cấp sửa mẫu màu theo BM04.",
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("requested_on", "Ngày yêu cầu", "date", required=True),
            _f("expected_colours", "Màu theo BM04"),
            _f("requirement", "Yêu cầu chỉnh sửa", required=True),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/16: one item per finding of the proof check.
        "template_id": "supply_chain.design_revision_request",
        "title": "YÊU CẦU SỬA THIẾT KẾ BAO BÌ",
        "doc_type": "design_revision_request",
        "description": (
            "Bước 12: Cung ứng yêu cầu sửa bản in thiết kế; mỗi điều hệ thống tìm thấy khi so"
            " bản in với BM04 và luật nhãn là một mục."
        ),
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("product_name", "Tên sản phẩm", required=True),
            _f("requested_on", "Ngày yêu cầu", "date", required=True),
            _table(
                "items",
                "Hạng mục cần sửa",
                [
                    ("subject", "Nội dung", "text"),
                    ("finding", "Hiện trạng", "text"),
                    ("requirement", "Yêu cầu", "text"),
                ],
                required=True,
            ),
            _f("notes", "Ghi chú"),
        ],
    },
    {
        # Ticket ai-automation/17: QC failed by the numbers; one item per
        # defect the report describes, the requirement a person writes.
        "template_id": "supply_chain.rework_request",
        "title": "YÊU CẦU LÀM LẠI SAU KIỂM HÀNG",
        "doc_type": "rework_request",
        "description": (
            "Bước 14: QC không đạt; tóm tắt số lỗi so với số chấp nhận (do hệ thống chép từ"
            " báo cáo QC) và từng lỗi cần nhà cung cấp làm lại."
        ),
        "fields": [
            _f("po_reference", "Số PO", required=True),
            _f("supplier_name", "Nhà cung cấp", required=True),
            _f("requested_on", "Ngày yêu cầu", "date", required=True),
            _f("qc_summary", "Kết quả kiểm (số lỗi / số chấp nhận)"),
            _table(
                "items",
                "Lỗi cần làm lại",
                [("defect", "Lỗi", "text"), ("requirement", "Yêu cầu", "text")],
                required=True,
            ),
            _f("notes", "Ghi chú"),
        ],
    },
]

VERSION = "1.0.0"


def _docx(spec: dict[str, Any]) -> DocxDocument:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(spec["title"])
    run.bold = True
    run.font.size = Pt(15)
    scalars = [f for f in spec["fields"] if f.get("kind") != "table"]
    tables = [f for f in spec["fields"] if f.get("kind") == "table"]
    info = document.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for f in scalars:
        if f["name"] == "notes":
            continue
        cells = info.add_row().cells
        cells[0].text = f["label"]
        cells[1].text = "{{" + f["name"] + "}}"
    for f in tables:
        document.add_paragraph()
        heading = document.add_paragraph()
        heading.add_run(f["label"]).bold = True
        columns = f["columns"]
        table = document.add_table(rows=2, cols=len(columns))
        table.style = "Table Grid"
        for index, column in enumerate(columns):
            table.cell(0, index).text = column["label"]
            table.cell(1, index).text = "{{" + f["name"] + "." + column["name"] + "}}"
    if any(f["name"] == "notes" for f in spec["fields"]):
        document.add_paragraph()
        document.add_paragraph("Ghi chú: {{notes}}")
    document.add_paragraph()
    signatures = document.add_table(rows=1, cols=2)
    signatures.cell(0, 0).text = "Người lập\n(Ký, ghi rõ họ tên)"
    signatures.cell(0, 1).text = "Người duyệt\n(Ký, ghi rõ họ tên)"
    return document


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for spec in TEMPLATES:
        version = spec.get("version", VERSION)
        ref = f"{spec['template_id']}@{version}"
        yaml_path, docx_path = OUT / f"{ref}.yaml", OUT / f"{ref}.docx"
        if yaml_path.exists() or docx_path.exists():
            print(f"exists, left as released: {ref}")
            continue
        declaration = {
            "schema_version": "1.0",
            "template_id": spec["template_id"],
            "version": version,
            "title": spec["title"],
            "doc_type": spec["doc_type"],
            "description": spec["description"],
            "fields": spec["fields"],
        }
        yaml_path.write_text(
            yaml.safe_dump(declaration, allow_unicode=True, sort_keys=False), encoding="utf-8"
        )
        _docx(spec).save(str(docx_path))
        print(f"wrote {ref}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
